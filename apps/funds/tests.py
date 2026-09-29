"""
Tests — apps.funds
==================
Bao phủ 10 trường hợp bắt buộc của đặc tả (Task 3-b):
    1. THU cập nhật số dư + mã phiếu PT.
    2. CHI cập nhật số dư + mã phiếu PC.
    3. CHI vượt số dư → InsufficientFundException (HTTP 400 envelope).
    4. so_tien <= 0 → 400.
    5. Bất biến I1: Σ thu − Σ chi = so_du_sau cuối (verify_balance_invariant).
    6. Khóa sổ kỳ → giao dịch mới trong kỳ bị PeriodLockedException (409).
    7. MEMBER gọi POST /funds/ → 403.
    8. GET /stats/ đúng tổng + low_balance flag.
    9. export_excel 200 + content-type xlsx.
    10. 2 giao dịch cùng loại cùng ngày → mã phiếu khác nhau (seq tăng).

Môi trường: settings.TESTING=True → axes tắt; throttle vẫn bật (burst 10/s)
nên `cache.clear()` trong setUp để reset lịch sử throttle giữa các test.
"""
from datetime import timedelta
from unittest import mock
from io import BytesIO
from typing import Optional

from django.core.cache import cache
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase
from django.db import connection
from django.test import TransactionTestCase
from unittest import skipIf

from apps.authentication.models import User
from apps.common.exceptions import (
    DuplicateDataException,
    InsufficientFundException,
    PeriodLockedException,
    ValidationException,
)
from apps.funds.models import FundPeriodLock, FundTransaction
from apps.funds.services import FundService, TransactionFactory
from apps.members.models import MemberProfile
from apps.funds.views import FundTransactionListCreateView
from core.algorithms.fund_invariants import FundInvariantsEngine


def make_user(email: str, role: str = "MEMBER") -> User:
    """Tạo user theo quy ước môi trường test (username sinh từ email — unique)."""
    return User.objects.create_user(
        username=email.split("@")[0],
        email=email,
        password="TestPass123!",
        role=role,
    )


class FundServiceTests(APITestCase):
    """Nhóm test nghiệp vụ Sổ quỹ ở tầng Service (không qua HTTP)."""

    def setUp(self) -> None:
        cache.clear()  # reset throttle history giữa các test
        self.bcn: User = make_user("bcn@clbip.test", role="BCN")

    # ------------------------------------------------------------------
    # (1) THU tiền
    # ------------------------------------------------------------------
    def test_thu_tien_cap_nhat_so_du_va_ma_phieu_PT(self) -> None:
        """THU 1.000.000 → so_du_sau=1.000.000, ma_phieu bắt đầu 'PT'."""
        tx = FundService.execute_transaction(
            loai_gd="THU",
            so_tien=1_000_000,
            nguoi_thuc_hien="Nguyễn Văn A",
        )
        self.assertEqual(tx.so_tien, 1_000_000)
        self.assertEqual(tx.so_du_sau, 1_000_000)
        self.assertTrue(tx.ma_phieu.startswith("PT"), f"Mã phiếu sai: {tx.ma_phieu}")
        self.assertFalse(tx.is_locked)
        # Đúng format PT + YYYYMMDD + 3 chữ số
        today = timezone.localdate()
        self.assertEqual(tx.ma_phieu, f"PT{today:%Y%m%d}001")

    # ------------------------------------------------------------------
    # (2) CHI tiền
    # ------------------------------------------------------------------
    def test_chi_tien_cap_nhat_so_du_va_ma_phieu_PC(self) -> None:
        """Thu 1.000.000 rồi CHI 300.000 → so_du_sau=700.000, mã phiếu 'PC'."""
        FundService.execute_transaction("THU", 1_000_000, "Nguyễn Văn A")
        tx = FundService.execute_transaction("CHI", 300_000, "Trần Thị B")
        self.assertEqual(tx.so_du_sau, 700_000)
        self.assertTrue(tx.ma_phieu.startswith("PC"), f"Mã phiếu sai: {tx.ma_phieu}")

    # ------------------------------------------------------------------
    # (3) CHI vượt số dư
    # ------------------------------------------------------------------
    def test_chi_vuot_so_du_raise_insufficient_fund(self) -> None:
        """Quỹ rỗng mà CHI → InsufficientFundException (không ghi sổ)."""
        before = FundTransaction.objects.count()
        with self.assertRaises(InsufficientFundException):
            FundService.execute_transaction("CHI", 500_000, "Trần Thị B")
        self.assertEqual(FundTransaction.objects.count(), before, "Không được ghi sổ khi lỗi")

    # ------------------------------------------------------------------
    # (4) Dữ liệu không hợp lệ
    # ------------------------------------------------------------------
    def test_loai_gd_va_so_tien_khong_hop_le(self) -> None:
        """loai_gd lạ / so_tien <= 0 đều bị chặn ở Service (400)."""
        with self.assertRaises(ValidationException):
            TransactionFactory.create("RUT_VO", so_tien=1000, ngay_gd=timezone.localdate())
        with self.assertRaises(ValidationException):
            FundService.execute_transaction("THU", 0, "A")
        with self.assertRaises(ValidationException):
            FundService.execute_transaction("CHI", -50_000, "A")

    # ------------------------------------------------------------------
    # (5) Bất biến I1 sau 3 giao dịch
    # ------------------------------------------------------------------
    def test_bat_bien_so_du_sau_3_giao_dich(self) -> None:
        """Sau 3 giao dịch: Σ thu − Σ chi == so_du_sau cuối (bất biến I1)."""
        FundService.execute_transaction("THU", 1_000_000, "A")
        FundService.execute_transaction("CHI", 300_000, "B")
        FundService.execute_transaction("THU", 200_000, "A")

        rows = list(
            FundTransaction.objects.order_by("ngay_gd", "id").values(
                "loai_gd", "so_tien", "so_du_sau"
            )
        )
        self.assertEqual(len(rows), 3)
        self.assertTrue(
            FundInvariantsEngine.verify_balance_invariant(rows),
            "Bất biến I1 bị phá vỡ!",
        )
        final_balance = rows[-1]["so_du_sau"]
        expected = 1_000_000 - 300_000 + 200_000
        self.assertEqual(final_balance, expected)
        self.assertTrue(FundInvariantsEngine.verify_balance_chain(rows))

    # ------------------------------------------------------------------
    # (6) Khóa sổ kỳ
    # ------------------------------------------------------------------
    def test_ky_khoa_chan_giao_dich_moi(self) -> None:
        """Khóa kỳ hôm nay → giao dịch mới trong kỳ bị PeriodLockedException."""
        FundService.execute_transaction("THU", 1_000_000, "A")
        today = timezone.localdate()
        lock = FundService.lock_period(
            ten_ky="Kỳ 09/2026",
            tu_ngay=today,
            den_ngay=today,
            locked_by=self.bcn,
        )
        self.assertIsNotNone(lock.pk)
        # Giao dịch cũ trong kỳ bị đóng băng
        self.assertTrue(FundTransaction.objects.filter(is_locked=True).exists())

        with self.assertRaises(PeriodLockedException):
            FundService.execute_transaction("CHI", 100_000, "B", ngay_gd=timezone.now())
        # Ngoài kỳ thì vẫn ghi được
        tx = FundService.execute_transaction(
            "THU", 50_000, "B", ngay_gd=timezone.now() + timedelta(days=40)
        )
        self.assertFalse(tx.is_locked)

    def test_lock_period_trung_ten_ky(self) -> None:
        """ten_ky trùng → DuplicateDataException (409); tu_ngay > den_ngay → 400."""
        today = timezone.localdate()
        FundService.lock_period("Kỳ duy nhất", today, today, self.bcn)
        with self.assertRaises(DuplicateDataException):
            FundService.lock_period("Kỳ duy nhất", today, today, self.bcn)
        with self.assertRaises(ValidationException):
            FundService.lock_period("Kỳ sai", today, today - timedelta(days=1), self.bcn)

    # ------------------------------------------------------------------
    # (8) Thống kê + low_balance (service-level bổ trợ)
    # ------------------------------------------------------------------
    def test_stats_service_level_low_balance(self) -> None:
        """balance=100.000 < 200.000 → low_balance=True; thu thêm → False."""
        FundService.execute_transaction("THU", 1_000_000, "A")
        FundService.execute_transaction("CHI", 900_000, "B")
        stats = FundService.get_stats()
        self.assertEqual(stats["total_income"], 1_000_000)
        self.assertEqual(stats["total_expense"], 900_000)
        self.assertEqual(stats["balance"], 100_000)
        self.assertTrue(stats["low_balance"])

        FundService.execute_transaction("THU", 500_000, "A")
        stats = FundService.get_stats()
        self.assertFalse(stats["low_balance"], "Balance 600k phải trên ngưỡng 200k")


class FundAPITests(APITestCase):
    """Nhóm test qua HTTP: envelope, phân quyền, mã lỗi, file Excel."""

    def setUp(self) -> None:
        cache.clear()
        self.bcn: User = make_user("bcn.api@clbip.test", role="BCN")
        self.member: User = make_user("member.api@clbip.test", role="MEMBER")
        self.client = APIClient()
        self.client.force_authenticate(user=self.bcn)
        self.list_url = reverse("fund_list")
        self.stats_url = reverse("fund_stats")
        self.lock_url = reverse("fund_lock_period")
        self.locks_url = reverse("fund_locks")
        self.export_url = reverse("fund_export_excel")

    def _post_transaction(self, payload: dict):
        """Tiện ích POST /api/v1/funds/."""
        return self.client.post(self.list_url, payload, format="json")

    # ------------------------------------------------------------------
    # (3) CHI vượt số dư qua API → 400 envelope success=false
    # ------------------------------------------------------------------
    def test_api_chi_vuot_so_du_tra_400_envelope(self) -> None:
        resp = self._post_transaction(
            {"loai_gd": "CHI", "so_tien": 500_000, "nguoi_thuc_hien": "Trần Thị B"}
        )
        self.assertEqual(resp.status_code, 400)
        body = resp.json()
        self.assertFalse(body["success"])
        # Message lấy từ FundInvariantsEngine (bất biến I3: chi không được làm dư âm)
        self.assertIn("số dư quỹ", body["message"])
        self.assertIsNone(body["data"])

    # ------------------------------------------------------------------
    # (4) so_tien <= 0 qua API → 400
    # ------------------------------------------------------------------
    def test_api_so_tien_khong_hop_le_tra_400(self) -> None:
        for so_tien in (0, -1000):
            with self.subTest(so_tien=so_tien):
                resp = self._post_transaction(
                    {"loai_gd": "THU", "so_tien": so_tien, "nguoi_thuc_hien": "A"}
                )
                self.assertEqual(resp.status_code, 400)
                self.assertFalse(resp.json()["success"])

    # ------------------------------------------------------------------
    # POST hợp lệ qua API → 201 + envelope
    # ------------------------------------------------------------------
    def test_api_lap_phieu_thanh_cong_201(self) -> None:
        resp = self._post_transaction(
            {"loai_gd": "THU", "so_tien": 1_000_000, "nguoi_thuc_hien": "Nguyễn Văn A"}
        )
        self.assertEqual(resp.status_code, 201)
        body = resp.json()
        self.assertTrue(body["success"])
        self.assertEqual(body["message"], "Lập phiếu thành công")
        self.assertEqual(body["data"]["so_du_sau"], 1_000_000)
        self.assertTrue(body["data"]["ma_phieu"].startswith("PT"))

    # ------------------------------------------------------------------
    # (6b) Khóa kỳ qua API → 409
    # ------------------------------------------------------------------
    def test_api_giao_dich_trong_ky_khoa_tra_409(self) -> None:
        today = timezone.localdate().isoformat()
        lock_resp = self.client.post(
            self.lock_url,
            {"ten_ky": "Kỳ API 09", "tu_ngay": today, "den_ngay": today},
            format="json",
        )
        self.assertEqual(lock_resp.status_code, 200)
        self.assertEqual(lock_resp.json()["message"], "Khóa sổ kỳ thành công")

        resp = self._post_transaction(
            {"loai_gd": "CHI", "so_tien": 100_000, "nguoi_thuc_hien": "B"}
        )
        self.assertEqual(resp.status_code, 409)
        self.assertFalse(resp.json()["success"])

    # ------------------------------------------------------------------
    # (7) MEMBER bị chặn 403
    # ------------------------------------------------------------------
    def test_api_member_post_bi_cho_403(self) -> None:
        self.client.force_authenticate(user=self.member)
        resp = self._post_transaction(
            {"loai_gd": "THU", "so_tien": 100_000, "nguoi_thuc_hien": "Hacker"}
        )
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(resp.json()["success"])
        self.assertEqual(FundTransaction.objects.count(), 0)

    def test_api_chua_dang_nhap_401(self) -> None:
        client = APIClient()  # không authenticate
        resp = client.get(self.list_url)
        self.assertEqual(resp.status_code, 401)

    # ------------------------------------------------------------------
    # (8) GET /stats/
    # ------------------------------------------------------------------
    def test_api_stats_dung_va_low_balance(self) -> None:
        FundService.execute_transaction("THU", 100_000, "A")
        resp = self.client.get(self.stats_url)
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertTrue(body["success"])
        data = body["data"]
        self.assertEqual(data["total_income"], 100_000)
        self.assertEqual(data["total_expense"], 0)
        self.assertEqual(data["balance"], 100_000)
        self.assertTrue(data["low_balance"], "100k < 200k phải cảnh báo quỹ thấp")

    # ------------------------------------------------------------------
    # (9) export_excel
    # ------------------------------------------------------------------
    def test_api_export_excel_200_va_content_type(self) -> None:
        FundService.execute_transaction("THU", 1_000_000, "A")
        FundService.execute_transaction("CHI", 250_000, "B")
        resp = self.client.get(self.export_url)
        self.assertEqual(resp.status_code, 200)
        expected_ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        self.assertEqual(resp["Content-Type"], expected_ct)
        self.assertIn("so_quy_clbip.xlsx", resp["Content-Disposition"])

        # Kiểm chứng nội dung workbook
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["Sổ quỹ"]
        self.assertEqual(ws.max_row, 3, "1 header + 2 giao dịch")
        self.assertEqual([c.value for c in ws[1]][0], "Mã phiếu")
        self.assertEqual(ws.cell(row=2, column=3).value, "1,000,000")

    # ------------------------------------------------------------------
    # (10) Mã phiếu idempotent: cùng loại cùng ngày → seq tăng
    # ------------------------------------------------------------------
    def test_api_hai_giao_dich_cung_loai_ma_phieu_khac_nhau(self) -> None:
        r1 = self._post_transaction(
            {"loai_gd": "THU", "so_tien": 100_000, "nguoi_thuc_hien": "A"}
        )
        r2 = self._post_transaction(
            {"loai_gd": "THU", "so_tien": 200_000, "nguoi_thuc_hien": "A"}
        )
        self.assertEqual(r1.status_code, 201)
        self.assertEqual(r2.status_code, 201)
        ma1, ma2 = r1.json()["data"]["ma_phieu"], r2.json()["data"]["ma_phieu"]
        self.assertNotEqual(ma1, ma2)
        today = timezone.localdate()
        self.assertEqual(ma1, f"PT{today:%Y%m%d}001")
        self.assertEqual(ma2, f"PT{today:%Y%m%d}002")

    # ------------------------------------------------------------------
    # GET list: lọc + sort whitelist + envelope pagination
    # ------------------------------------------------------------------
    def test_api_danh_sach_phan_trang_va_loc(self) -> None:
        FundService.execute_transaction("THU", 1_000_000, "A")
        FundService.execute_transaction("CHI", 200_000, "B")

        resp = self.client.get(self.list_url)
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertTrue(body["success"])
        self.assertEqual(body["data"]["pagination"]["total_items"], 2)
        self.assertEqual(len(body["data"]["items"]), 2)

        resp = self.client.get(self.list_url, {"loai_gd": "CHI"})
        items = resp.json()["data"]["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["loai_gd"], "CHI")

        resp = self.client.get(self.list_url, {"sort": "so_tien; DROP TABLE x"})
        self.assertEqual(resp.status_code, 200, "Sort ngoài whitelist phải fallback an toàn")

    def test_api_locks_danh_sach_ky(self) -> None:
        today = timezone.localdate()
        FundService.lock_period("Kỳ liệt kê", today, today, self.bcn)
        resp = self.client.get(self.locks_url)
        self.assertEqual(resp.status_code, 200)
        items = resp.json()["data"]["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["ten_ky"], "Kỳ liệt kê")


class FundIdempotencyTests(APITestCase):
    """
    QA-Audit nhóm 3 — Idempotency-Key chống ghi trùng sổ quỹ:
    - Submit 2 lần cùng key → 1 giao dịch duy nhất, lần 2 trả HTTP 200 + phiếu cũ.
    - Key khác/giống → hành vi đúng.
    - Header "Idempotency-Key" và body field đều nhận.
    """

    def setUp(self) -> None:
        cache.clear()
        self.bcn = User.objects.create_user(
            email="idem-bcn@clbip.test", password="TestPass123!", role="BCN",
        )
        MemberProfile.objects.create(user=self.bcn, ho_ten="Thủ Quỹ Idem")
        self.client.force_authenticate(user=self.bcn)
        self.url = reverse("fund_list")
        self.payload = {
            "loai_gd": "THU",
            "so_tien": 500_000,
            "nguoi_thuc_hien": "Nhà tài trợ XYZ",
            "hinh_thuc": "CHUYEN_KHOAN",
        }

    def test_01_double_submit_cung_key_chi_tao_mot_giao_dich(self) -> None:
        res1 = self.client.post(
            self.url, {**self.payload, "idempotency_key": "hop-dong-abc-001"}, format="json",
        )
        self.assertEqual(res1.status_code, 201, res1.data)
        res2 = self.client.post(
            self.url, {**self.payload, "idempotency_key": "hop-dong-abc-001"}, format="json",
        )
        # Lần 2: HTTP 200 + trả lại giao dịch cũ (không phải 201)
        self.assertEqual(res2.status_code, 200, res2.data)
        self.assertIn("đã tồn tại", res2.data["message"])
        self.assertEqual(
            res1.data["data"]["ma_phieu"], res2.data["data"]["ma_phieu"],
        )
        self.assertEqual(FundTransaction.objects.count(), 1)

    def test_02_header_idempotency_key_uu_tien(self) -> None:
        res1 = self.client.post(
            self.url, self.payload, format="json", HTTP_IDEMPOTENCY_KEY="header-key-42",
        )
        self.assertEqual(res1.status_code, 201)
        # Gửi lại cùng header (body có key khác) — header thắng
        res2 = self.client.post(
            self.url, {**self.payload, "idempotency_key": "body-key-khac"},
            format="json", HTTP_IDEMPOTENCY_KEY="header-key-42",
        )
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(FundTransaction.objects.count(), 1)

    def test_03_khong_key_ghi_binh_thuong(self) -> None:
        res1 = self.client.post(self.url, self.payload, format="json")
        res2 = self.client.post(self.url, self.payload, format="json")
        self.assertEqual(res1.status_code, 201)
        self.assertEqual(res2.status_code, 201)
        self.assertEqual(FundTransaction.objects.count(), 2)

    def test_04_key_qua_dai_400(self) -> None:
        res = self.client.post(
            self.url, self.payload, format="json",
            HTTP_IDEMPOTENCY_KEY="x" * 65,
        )
        self.assertEqual(res.status_code, 400)


class FundIdempotencyRaceTests(TransactionTestCase):
    """
    Race test (QA-Audit nhóm 3): N thread POST song song cùng Idempotency-Key
    → CHÍNH XÁC 1 giao dịch được ghi (kiểm tra trong atomic sau select_for_update).
    Dùng TransactionTestCase vì TestCase bọc transaction — thread không thấy data.
    """

    @mock.patch.object(FundTransactionListCreateView, "throttle_classes", [])
    @skipIf(
        connection.vendor == "sqlite",
        "SQLite khóa toàn bảng khi 2 connection cùng ghi — race test chỉ chạy "
        "đúng ý trên MySQL/PostgreSQL (CSDL chính của dự án là MySQL).",
    )
    def test_04_thread_song_song_cung_key(self) -> None:
        """4 thread song song cùng Idempotency-Key → đúng 1 giao dịch (201), 3 replay (200)."""
        from threading import Barrier, Thread

        bcn = User.objects.create_user(
            email="race-bcn@clbip.test", password="TestPass123!", role="BCN",
        )
        MemberProfile.objects.create(user=bcn, ho_ten="Thủ Quỹ Race")
        barrier = Barrier(4)
        results: list = []

        def post_once() -> None:
            # Thread riêng → client riêng (force_authenticate gắn vào client)
            client = APIClient()
            client.force_authenticate(user=bcn)
            barrier.wait(timeout=10)  # đồng loạt bấm cùng lúc
            res = client.post(
                reverse("fund_list"),
                {
                    "loai_gd": "THU", "so_tien": 250_000,
                    "nguoi_thuc_hien": "Race Test",
                    "idempotency_key": "race-key-độc-nhất",
                },
                format="json",
            )
            results.append(res.status_code)

        threads = [Thread(target=post_once) for _ in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=30)

        # Chỉ 1 giao dịch ghi thành công với key này
        self.assertEqual(
            FundTransaction.objects.filter(idempotency_key="race-key-độc-nhất").count(), 1,
        )
        # Mỗi thread nhận đúng 1 kết quả: 201 (người thắng) hoặc 200 (replay)
        self.assertEqual(len(results), 4)
        self.assertTrue(all(code in (200, 201) for code in results), results)
        self.assertEqual(results.count(201), 1)
