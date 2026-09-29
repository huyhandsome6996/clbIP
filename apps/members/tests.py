"""
Tests — apps.members
CRUD phân quyền, IDOR protection, Trie search, Excel import/export, whitelist sort.
"""
import io
from typing import Any

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.authentication.models import User
from apps.members.models import MemberProfile


def make_user(email: str, role: str = "MEMBER", ho_ten: str = "") -> Any:
    """Helper: tạo User (+ MemberProfile nếu là thành viên)."""
    user = User.objects.create_user(email=email, password="TestPass123!", role=role)
    if ho_ten:
        MemberProfile.objects.create(user=user, ho_ten=ho_ten, lop="22A401")
    return user


class MemberCRUDPermissionTests(TestCase):
    """Phân quyền tạo/xem/sửa/khóa thành viên."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("bcn@clbip.vn", role="BCN", ho_ten="Chủ Nhiệm")
        self.member = make_user("member@clbip.vn", role="MEMBER", ho_ten="Thành Viên A")

    def test_bcn_create_member_success(self) -> None:
        """BCN tạo member → 201, DB có User + Profile."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.post(
            "/api/v1/members/",
            {"email": "new@clbip.vn", "mssv": "22A401999", "ho_ten": "Người Mới", "lop": "22A402"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertTrue(res.data["success"])
        self.assertTrue(User.objects.filter(email="new@clbip.vn").exists())
        profile = MemberProfile.objects.get(user__email="new@clbip.vn")
        self.assertEqual(profile.ho_ten, "Người Mới")

    def test_member_create_forbidden(self) -> None:
        """MEMBER tạo member → 403."""
        self.client.force_authenticate(user=self.member)
        res = self.client.post(
            "/api/v1/members/",
            {"email": "x@clbip.vn", "mssv": "22A401888", "ho_ten": "X"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_duplicate_email_rejected(self) -> None:
        """Email trùng → 400/409."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.post(
            "/api/v1/members/",
            {"email": "member@clbip.vn", "mssv": "22A401777", "ho_ten": "Trùng"},
            format="json",
        )
        self.assertIn(res.status_code, (status.HTTP_400_BAD_REQUEST, status.HTTP_409_CONFLICT))

    def test_member_cannot_view_other_profile360_idor(self) -> None:
        """⚠ IDOR: Member A xem profile360 của B → 403. Của chính mình → 200."""
        other = make_user("other@clbip.vn", role="MEMBER", ho_ten="Thành Viên B")
        self.client.force_authenticate(user=self.member)

        res_other = self.client.get(f"/api/v1/members/{other.member_profile.pk}/profile360/")
        self.assertEqual(res_other.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(res_other.data["success"])

        res_self = self.client.get(f"/api/v1/members/{self.member.member_profile.pk}/profile360/")
        self.assertEqual(res_self.status_code, status.HTTP_200_OK)
        self.assertEqual(res_self.data["data"]["ho_ten"], "Thành Viên A")

    def test_bcn_view_any_profile360(self) -> None:
        """BCN xem profile360 của bất kỳ ai → 200."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get(f"/api/v1/members/{self.member.member_profile.pk}/profile360/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["data"]["stats"]["xp_points"], 0)

    def test_member_cannot_edit_sensitive_fields(self) -> None:
        """Member tự sửa email/mssv/trạng thái → 400."""
        self.client.force_authenticate(user=self.member)
        res = self.client.patch(
            f"/api/v1/members/{self.member.member_profile.pk}/",
            {"mssv": "HACKED", "trang_thai_hd": "LEAVE"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_member_can_edit_own_basic_info(self) -> None:
        """Member tự sửa sdt của mình → 200."""
        self.client.force_authenticate(user=self.member)
        res = self.client.patch(
            f"/api/v1/members/{self.member.member_profile.pk}/",
            {"sdt": "0900900900"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.member.member_profile.refresh_from_db()
        self.assertEqual(self.member.member_profile.sdt, "0900900900")

    def test_lock_member(self) -> None:
        """BCN khóa member → is_active=False; member bị khóa không đăng nhập được."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.delete(f"/api/v1/members/{self.member.member_profile.pk}/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.member.refresh_from_db()
        self.assertFalse(self.member.is_active)

        # Đăng nhập bị từ chối
        client = APIClient()
        login = client.post(
            "/api/v1/auth/token/",
            {"email": "member@clbip.vn", "password": "TestPass123!"},
            format="json",
        )
        self.assertEqual(login.status_code, status.HTTP_401_UNAUTHORIZED)


class MemberSearchTests(TestCase):
    """Tìm kiếm Trie — hỗ trợ không dấu + MSSV."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("bcn2@clbip.vn", role="BCN", ho_ten="Ban Chủ Nhiệm")
        make_user("m1@clbip.vn", ho_ten="Nguyễn Nhật Nam")
        m2 = User.objects.create_user(email="m2@clbip.vn", password="TestPass123!")
        MemberProfile.objects.create(user=m2, ho_ten="Hoàng Bảo Trân", lop="22A402")
        m3 = User.objects.create_user(email="m3@clbip.vn", password="TestPass123!", mssv="22A401123")
        MemberProfile.objects.create(user=m3, ho_ten="Trần Văn Tèo", lop="22A403")

    def test_search_without_accents(self) -> None:
        """q='nguyen nhat' (không dấu) → tìm ra 'Nguyễn Nhật Nam'."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/search/?q=nguyen nhat")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        names = [item["ho_ten"] for item in res.data["data"]["items"]]
        self.assertIn("Nguyễn Nhật Nam", names)

    def test_search_by_mssv_prefix(self) -> None:
        """q='22A4011' → tìm theo MSSV."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/search/?q=22A4011")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        items = res.data["data"]["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["mssv"], "22A401123")

    def test_search_empty_query(self) -> None:
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/search/?q=")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["data"]["count"], 0)

    def test_search_requires_auth(self) -> None:
        res = APIClient().get("/api/v1/members/search/?q=nguyen")
        self.assertEqual(res.status_code, status.HTTP_401_UNAUTHORIZED)


class MemberListSortWhitelistTests(TestCase):
    """Whitelist sort — payload độc hại vẫn xử lý an toàn."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("bcn3@clbip.vn", role="BCN", ho_ten="BCN")
        make_user("s1@clbip.vn", ho_ten="Sinh Viên 1")
        make_user("s2@clbip.vn", ho_ten="Sinh Viên 2")

    def test_malicious_sort_param_is_sanitized(self) -> None:
        """?sort=email'; -- → không lỗi 500, fallback '-xp_points'."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/", {"sort": "email'; --"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.data["success"])

    def test_list_pagination_envelope(self) -> None:
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("pagination", res.data["data"])
        self.assertGreaterEqual(res.data["data"]["pagination"]["total_items"], 3)

    def test_member_list_requires_bcn(self) -> None:
        member = make_user("mm@clbip.vn", ho_ten="MM")
        self.client.force_authenticate(user=member)
        res = self.client.get("/api/v1/members/")
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)


class MemberExcelTests(TestCase):
    """Import / Export Excel."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("bcn4@clbip.vn", role="BCN", ho_ten="BCN Excel")

    def _xlsx_bytes(self, rows: list) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.append(["ho_ten", "email", "mssv", "lop", "sdt"])
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_import_excel_with_error_preview(self) -> None:
        """2 dòng hợp lệ + 1 dòng trùng mssv → created=2, errors=1."""
        self.client.force_authenticate(user=self.bcn)
        # Tạo trước user trùng MSSV
        make_user("exist@clbip.vn", ho_ten="Đã Tồn Tại")
        existing = User.objects.get(email="exist@clbip.vn")
        existing.mssv = "22A401111"
        existing.save()

        buf = self._xlsx_bytes([
            ["Sinh Viên 1", "sv1@clbip.vn", "22A401001", "22A401", "0901"],
            ["Sinh Viên 2", "sv2@clbip.vn", "22A401002", "22A401", "0902"],
            ["Trùng MSSV", "sv3@clbip.vn", "22A401111", "22A401", "0903"],
        ])
        upload = SimpleUploadedFile("thanh_vien.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res = self.client.post(
            "/api/v1/members/import-excel/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["data"]["created"], 2)
        self.assertEqual(len(res.data["data"]["errors"]), 1)
        self.assertIn("22A401111", res.data["data"]["errors"][0]["error"])

    def test_import_requires_xlsx(self) -> None:
        self.client.force_authenticate(user=self.bcn)
        upload = SimpleUploadedFile("not_excel.txt", b"not excel", content_type="text/plain")
        res = self.client.post(
            "/api/v1/members/import-excel/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_export_excel(self) -> None:
        """BCN xuất file → 200 + đọc lại được nội dung xlsx."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/export-excel/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", res["Content-Type"])
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "Thành viên")
        self.assertGreaterEqual(ws.max_row, 2)  # header + ít nhất 1 dòng


class Profile360DataTests(TestCase):
    """Dữ liệu tổng hợp hồ sơ 360°."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.member = make_user("p@clbip.vn", ho_ten="Profile Test")

    def test_profile360_contains_all_sections(self) -> None:
        self.client.force_authenticate(user=self.member)
        res = self.client.get(f"/api/v1/members/{self.member.member_profile.pk}/profile360/")
        data = res.data["data"]
        self.assertIn("stats", data)
        self.assertIn("badges", data)
        self.assertIn("board_positions", data)
        self.assertIn("recent_xp", data)
        self.assertIsNone(data["stats"]["attendance_rate_percent"])  # chưa có record

    def test_board_list_and_create(self) -> None:
        """GET /board/ mọi user; POST /board/ chỉ BCN."""
        from apps.members.models import BoardMember

        self.client.force_authenticate(user=self.member)
        res = self.client.get("/api/v1/members/board/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        res_post = self.client.post(
            "/api/v1/members/board/",
            {"member": self.member.member_profile.pk, "nhiem_ky": "2025-2026",
             "chuc_vu": "CHU_NHIEM", "ban_phu_trach": "HOC_THUAT"},
            format="json",
        )
        self.assertEqual(res_post.status_code, status.HTTP_403_FORBIDDEN)

        self.client.force_authenticate(user=make_user("b5@clbip.vn", role="BCN"))
        res_post2 = self.client.post(
            "/api/v1/members/board/",
            {"member": self.member.member_profile.pk, "nhiem_ky": "2025-2026",
             "chuc_vu": "CHU_NHIEM", "ban_phu_trach": "HOC_THUAT"},
            format="json",
        )
        self.assertEqual(res_post2.status_code, status.HTTP_201_CREATED)
        self.assertTrue(BoardMember.objects.filter(member=self.member.member_profile).exists())


class MemberSearchPIITests(TestCase):
    """
    QA-Audit 2a — chống rò rỉ PII qua tìm kiếm thành viên:
    - MEMBER thường: chỉ nhận trường công khai (id/ho_ten/lop/avatar/level/xp),
      KHÔNG có email/sdt/mssv; chỉ thấy thành viên đang HOẠT ĐỘNG.
    - BCN/ADMIN: nhận serializer đầy đủ phục vụ quản lý.
    """

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("piibcn@clbip.vn", role="BCN", ho_ten="Chủ Nhiệm Test")
        self.member = make_user("piimember@clbip.vn", role="MEMBER", ho_ten="Thành Viên PII")
        self.other = make_user("piiother@clbip.vn", role="MEMBER", ho_ten="Nguyễn Bình Yên")

    def test_member_search_khong_chua_pii(self) -> None:
        """MEMBER search → items không có email/sdt/mssv."""
        self.client.force_authenticate(user=self.member)
        res = self.client.get("/api/v1/members/search/", {"q": "nguyen"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        items = res.data["data"]["items"]
        self.assertGreaterEqual(len(items), 1)
        for field in ("email", "sdt", "mssv"):
            self.assertNotIn(field, items[0])
        for field in ("id", "ho_ten", "lop", "xp_points", "current_level"):
            self.assertIn(field, items[0])

    def test_member_search_chi_thay_thanh_vien_active(self) -> None:
        """MEMBER không thấy thành viên INACTIVE; BCN vẫn thấy."""
        profile = MemberProfile.objects.get(user__email="piiother@clbip.vn")
        profile.trang_thai_hd = MemberProfile.TrangThai.INACTIVE
        profile.save(update_fields=["trang_thai_hd"])

        self.client.force_authenticate(user=self.member)
        res = self.client.get("/api/v1/members/search/", {"q": "nguyen"})
        self.assertEqual(res.data["data"]["items"], [])

        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/search/", {"q": "nguyen"})
        self.assertGreaterEqual(len(res.data["data"]["items"]), 1)

    def test_bcn_search_nhan_serializer_day_du(self) -> None:
        """BCN search → có email (mssv có thể None nếu không đặt)."""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/search/", {"q": "nguyen"})
        items = res.data["data"]["items"]
        self.assertGreaterEqual(len(items), 1)
        self.assertIn("email", items[0])
        self.assertIn("mssv", items[0])

    def test_board_list_member_khong_nhan_mssv(self) -> None:
        """GET /members/board/ với MEMBER → không lộ member_mssv (PII cán bộ)."""
        from apps.members.models import BoardMember

        BoardMember.objects.create(
            member=MemberProfile.objects.get(user=self.bcn),
            nhiem_ky="2025-2026", chuc_vu="CHU_NHIEM", ban_phu_trach="HOC_THUAT",
        )
        self.client.force_authenticate(user=self.member)
        res = self.client.get("/api/v1/members/board/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        items = res.data["data"]["items"]
        self.assertGreaterEqual(len(items), 1)
        self.assertNotIn("member_mssv", items[0])

        self.client.force_authenticate(user=self.bcn)
        res = self.client.get("/api/v1/members/board/")
        self.assertIn("member_mssv", res.data["data"]["items"][0])


class MemberTrieCacheTests(TestCase):
    """
    QA-Audit nhóm 5 — Trie tìm kiếm member cache theo process:
    - Dùng lại instance giữa các lần tra cứu (không dựng lại mỗi request).
    - Tín hiệu post_save/post_delete (tạo/sửa/xóa MemberProfile, User mssv)
      vô hiệu hóa → lần tra cứu kế tiếp dựng lại với dữ liệu mới.
    """

    def test_trie_duoc_dung_lai_giua_cac_lan_goi(self) -> None:
        from apps.members import search_index

        search_index.invalidate()  # reset trạng thái module (cô lập giữa test)
        t1 = search_index.get_trie()
        t2 = search_index.get_trie()
        self.assertIs(t1, t2)

    def test_tao_member_moi_vao_trie_sau_invalidate(self) -> None:
        from apps.members import search_index

        search_index.invalidate()
        old_trie = search_index.get_trie()
        # Tạo member → post_save signal tự gọi invalidate
        make_user("trie-moi@clbip.vn", role="MEMBER", ho_ten="Triệu Test Mới")
        new_trie = search_index.get_trie()
        self.assertIsNot(old_trie, new_trie)
        # Tìm được thành viên mới qua tiền tố (không dấu)
        matched = new_trie.search_prefix("trieu test")
        self.assertGreaterEqual(len(matched), 1)

    def test_search_service_dung_trie_cache(self) -> None:
        """Service trả kết quả đúng khi dùng Trie cache (không regression)."""
        from apps.members import search_index
        from apps.members.services import MemberService

        search_index.invalidate()
        # "Thành Viên PII" (setUp của class khác) không tồn tại ở đây — tự tạo
        make_user("trie-svc@clbip.vn", role="MEMBER", ho_ten="Thanh Vien Trie")
        results = MemberService.search_profiles("thanh vien trie", limit=5)
        self.assertGreaterEqual(len(results), 1)


class MemberPasswordPolicyTests(TestCase):
    """QA-Audit P0/DoD — không còn mật khẩu mặc định công khai khi BCN tạo member."""

    def setUp(self) -> None:
        self.client = APIClient()
        self.bcn = make_user("pwd-bcn@clbip.vn", role="BCN", ho_ten="BCN Mật Khẩu")

    def test_tao_member_khong_password_nhan_mat_khau_ngau_nhien(self) -> None:
        """"""
        self.client.force_authenticate(user=self.bcn)
        res = self.client.post(
            "/api/v1/members/",
            {"email": "ngau-nhien@clbip.vn", "mssv": "22A401777", "ho_ten": "Mật Khẩu Ngẫu Nhiên"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertIn("Mật khẩu khởi tạo", res.data["message"])
        self.assertNotIn("CLBIP@2026", res.data["message"])
        # Mật khẩu in ra phải dùng được để đăng nhập (không phải mặc định công khai)
        password = res.data["message"].split("Mật khẩu khởi tạo: ")[1].split(" ")[0]
        login = self.client.post(
            "/api/v1/auth/token/", {"email": "ngau-nhien@clbip.vn", "password": password}, format="json",
        )
        self.assertEqual(login.status_code, status.HTTP_200_OK)


class MemberTrieCacheSharedVersionTests(TestCase):
    """
    QA-Audit đợt 2 (P2-2) — Trie thành viên phải invalidate ĐA WORKER:
    counter lưu trong cache dùng chung; bump counter ở "worker khác" →
    process này phải dựng lại Trie (trước đây chỉ tăng biến local).
    """

    def setUp(self) -> None:
        from django.core.cache import cache

        cache.clear()

    def test_bump_counter_worker_khac_buoc_rebuild(self) -> None:
        from django.core.cache import cache
        from django.db.models.signals import post_delete, post_save

        from apps.members import search_index

        user1 = User.objects.create_user(
            email="trie1@clbip.vn", password="TestPass123!", role="MEMBER", mssv="22A4010101"
        )
        MemberProfile.objects.create(user=user1, ho_ten="Nguyễn Trie Một")
        trie_1 = search_index.get_trie()
        self.assertIn(user1.member_profile.pk, trie_1.search_prefix("nguyen trie"))

        # Tạo hồ sơ KHÔNG bắn signal cục bộ (mô phỏng: worker khác nhận ghi)
        # ⚠️ disconnect phải truyền đúng sender — lookup key = (dispatch_uid, _make_id(sender))
        post_save.disconnect(sender=MemberProfile, dispatch_uid="trie_member_save")
        post_delete.disconnect(sender=MemberProfile, dispatch_uid="trie_member_del")
        post_save.disconnect(sender=User, dispatch_uid="trie_user_save")
        post_delete.disconnect(sender=User, dispatch_uid="trie_user_del")
        try:
            user2 = User.objects.create_user(
                email="trie2@clbip.vn", password="TestPass123!", role="MEMBER", mssv="22A4010102"
            )
            MemberProfile.objects.create(user=user2, ho_ten="Nguyễn Trie Hai")
        finally:
            post_save.connect(search_index.invalidate, sender=MemberProfile, dispatch_uid="trie_member_save")
            post_delete.connect(search_index.invalidate, sender=MemberProfile, dispatch_uid="trie_member_del")
            post_save.connect(search_index.invalidate, sender=User, dispatch_uid="trie_user_save")
            post_delete.connect(search_index.invalidate, sender=User, dispatch_uid="trie_user_del")

        # Chưa biết gì về thay đổi → Trie cũ giữ nguyên (không rebuild thừa)
        trie_stale = search_index.get_trie()
        self.assertIs(trie_stale, trie_1)
        self.assertNotIn(user2.member_profile.pk, trie_stale.search_prefix("nguyen trie"))

        # "Worker khác" tăng counter dùng chung → process này phải thấy
        try:
            cache.incr(search_index._VERSION_CACHE_KEY)
        except ValueError:
            cache.add(search_index._VERSION_CACHE_KEY, 1, timeout=None)

        trie_2 = search_index.get_trie()
        self.assertIsNot(trie_2, trie_stale, "Counter dùng chung đổi → phải dựng lại Trie")
        self.assertIn(user2.member_profile.pk, trie_2.search_prefix("nguyen trie"))
